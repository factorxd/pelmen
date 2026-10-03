import os
import tempfile
import zipfile
import shutil

import pandas as pd
from docx import Document
from docxcompose.composer import Composer
from PySide6.QtCore import QObject, Signal

from logic.doc_generator import generate_docx


class MassGenerateWorker(QObject):
    """Массовая генерация в отдельном потоке."""

    progress = Signal(int, int)   # current, total
    finished = Signal(str, int)   # message, count
    error = Signal(str, int)      # message, row_index (-1 = не привязано к строке)
    phase_changed = Signal(str)  # "generating" | "merging" | "saving"
    cancelled = Signal()

    def __init__(self, template, df, mapping, mode, output_path,
                 name_mask="", get_field_info=None, format_value=None,
                 parent=None):
        super().__init__(parent)
        self.template = template
        self.df = df
        self.mapping = mapping
        self.mode = mode                # "single" | "multi_zip" | "multi_folder"
        self.output_path = output_path
        self.name_mask = name_mask
        self.get_field_info = get_field_info
        self.format_value = format_value
        self._cancelled = False

    def cancel(self):
        self._cancelled = True

    def run(self):
        try:
            if self.mode == "single":
                self._run_single()
            elif self.mode in ("multi_zip", "multi_folder"):
                self._run_multi()
            else:
                self.error.emit(f"Неизвестный режим: {self.mode}", -1)
        except Exception as e:
            self.error.emit(f"Непредвиденная ошибка: {e}", -1)

    def _build_data_dict(self, row):
        data = {}
        for col, (typ, field_name) in self.mapping.items():
            value = row[col]
            if pd.isna(value):
                value = ""
            info = self.get_field_info(field_name) if self.get_field_info else {}
            if self.format_value:
                data[field_name] = self.format_value(value, info)
            else:
                data[field_name] = str(value)
        for block in self.template.blocks:
            data.setdefault(block.name, [])
        return data

    def _run_single(self):
        total = len(self.df)
        total_units = total * 2  # генерация + объединение
        temp_dir = tempfile.mkdtemp(prefix="pelmen_merge_")
        try:
            # Фаза 1: генерация (0 .. total)
            self.phase_changed.emit("generating")
            doc_paths = []
            for i, (_, row) in enumerate(self.df.iterrows(), start=1):
                if self._cancelled:
                    self.cancelled.emit()
                    return
                data = self._build_data_dict(row)
                part = os.path.join(temp_dir, f"part_{i:05d}.docx")
                try:
                    generate_docx(self.template.file_path, data, part)
                except Exception as e:
                    self.error.emit(f"Ошибка в строке {i}: {e}", i - 1)
                    return
                if i < total:
                    doc = Document(part)
                    doc.add_page_break()
                    doc.save(part)
                doc_paths.append(part)
                self.progress.emit(i, total_units)

            if not doc_paths:
                self.error.emit("Нет данных для генерации", -1)
                return

            # Фаза 2: объединение (total .. 2*total)
            self.phase_changed.emit("merging")
            composer = Composer(Document(doc_paths[0]))
            self.progress.emit(total + 1, total_units)
            for idx, path in enumerate(doc_paths[1:], start=2):
                if self._cancelled:
                    self.cancelled.emit()
                    return
                composer.append(Document(path))
                self.progress.emit(total + idx, total_units)

            self.phase_changed.emit("saving")
            composer.save(self.output_path)
            self.progress.emit(total_units, total_units)
            self.finished.emit(
                f"Сохранено {total} документов:\n{self.output_path}", total)
        finally:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def _run_multi(self):
        total = len(self.df)
        total_units = total + 1  # генерация + упаковка/перемещение
        with tempfile.TemporaryDirectory(prefix="pelmen_multi_") as tmpdir:
            self.phase_changed.emit("generating")
            outputs = []
            for i, (_, row) in enumerate(self.df.iterrows(), start=1):
                if self._cancelled:
                    self.cancelled.emit()
                    return
                data = self._build_data_dict(row)
                fname = self._make_filename(data, i)
                out = os.path.join(tmpdir, f"{fname}.docx")
                try:
                    generate_docx(self.template.file_path, data, out)
                except Exception as e:
                    self.error.emit(f"Ошибка в строке {i}: {e}", i - 1)
                    return
                outputs.append(out)
                self.progress.emit(i, total_units)

            self.phase_changed.emit("saving")
            if self.mode == "multi_zip":
                try:
                    with zipfile.ZipFile(self.output_path, "w",
                                         zipfile.ZIP_DEFLATED) as zf:
                        for path in outputs:
                            zf.write(path, os.path.basename(path))
                except Exception as e:
                    self.error.emit(f"Не удалось создать ZIP: {e}", -1)
                    return
                self.progress.emit(total_units, total_units)
                self.finished.emit(
                    f"Создано {total} документов.\nАрхив: {self.output_path}",
                    total)
            else:
                moved = 0
                for path in outputs:
                    dest = os.path.join(self.output_path, os.path.basename(path))
                    try:
                        shutil.move(path, dest)
                        moved += 1
                    except Exception as e:
                        self.error.emit(
                            f"Не удалось сохранить {os.path.basename(path)}: {e}",
                            -1)
                        return
                self.progress.emit(total_units, total_units)
                self.finished.emit(
                    f"Сохранено {moved} документов в папку:\n{self.output_path}",
                    moved)

    def _make_filename(self, data, idx):
        if self.name_mask:
            fname = self.name_mask
            for key, val in data.items():
                fname = fname.replace(f"{{{key}}}", str(val))
            fname = "".join(c for c in fname if c.isalnum() or c in "._- ")
            if fname:
                return fname
        return f"doc_{idx}"